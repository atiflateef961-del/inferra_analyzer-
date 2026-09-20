from __future__ import annotations

import csv
import io
import json
import logging
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader
from docx import Document as DocxDocument

ALLOWED_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".pdf",
    ".docx",
    ".txt",
    ".json",
    ".png",
    ".jpg",
    ".jpeg",
}

logger = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}

REVENUE_KEYS = {"revenue", "sales", "income", "amount", "total", "turnover"}
EXPENSE_KEYS = {
    "expense",
    "expenses",
    "cost",
    "costs",
    "cogs",
    "spend",
    "salary",
    "salaries",
    "wage",
    "wages",
    "payroll",
    "bonus",
    "bonuses",
    "overtime",
    "tax",
    "taxes",
}
PROFIT_KEYS = {"profit", "net", "margin", "earnings"}
CATEGORY_KEYS = {"category", "product", "sku", "item", "department", "channel"}
PERIOD_KEYS = {"month", "date", "period", "week", "day"}
LOCATION_KEYS = {
    "country",
    "city",
    "state",
    "province",
    "region",
    "location",
    "territory",
    "market",
    "continent",
}


def _normalize_header(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", value.strip().lower())


def _parse_number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "-"}:
        return None
    text = text.replace(",", "").replace("$", "").replace("%", "")
    try:
        return float(text)
    except ValueError:
        return None


def _classify_column(header: str) -> str | None:
    if header in PERIOD_KEYS or any(key in header for key in PERIOD_KEYS):
        return "period"
    if header in LOCATION_KEYS or any(key in header for key in LOCATION_KEYS):
        return "location"
    if header in CATEGORY_KEYS or any(key in header for key in CATEGORY_KEYS):
        return "category"
    if header in PROFIT_KEYS:
        return "profit"
    if header in EXPENSE_KEYS or any(key in header for key in EXPENSE_KEYS):
        return "expense"
    if header in REVENUE_KEYS or any(key in header for key in REVENUE_KEYS):
        return "revenue"
    return None


def _rows_from_csv(content: bytes) -> tuple[list[str], list[list[Any]]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [row for row in reader if any(str(cell).strip() for cell in row)]
    if not rows:
        return [], []
    return [str(cell) for cell in rows[0]], rows[1:]


def _rows_from_text_table(content: bytes) -> tuple[list[str], list[list[Any]]] | None:
    """Recognize text files that contain a consistent CSV/TSV-style table."""
    text = content.decode("utf-8-sig", errors="replace")
    sample_rows = [line for line in text.splitlines() if line.strip()][:6]
    if len(sample_rows) < 3:
        return None
    try:
        dialect = csv.Sniffer().sniff("\n".join(sample_rows), delimiters=",\t;|")
    except csv.Error:
        return None
    rows = [row for row in csv.reader(io.StringIO(text), dialect) if any(str(cell).strip() for cell in row)]
    if len(rows) < 3 or len(rows[0]) < 2 or any(len(row) != len(rows[0]) for row in rows[1:]):
        return None
    return [str(cell) for cell in rows[0]], rows[1:]


def _rows_from_xlsx(content: bytes) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    try:
        sheet = workbook.active
        rows = [
            [cell for cell in row]
            for row in sheet.iter_rows(values_only=True)
            if any(cell not in (None, "") for cell in row)
        ]
    finally:
        workbook.close()
    if not rows:
        return [], []
    return [str(cell) if cell is not None else "" for cell in rows[0]], rows[1:]


def _extract_pdf_text(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages).strip()


def _extract_docx_text(content: bytes) -> str:
    document = DocxDocument(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _extract_json_text(content: bytes) -> str:
    try:
        payload = json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise ValueError("The JSON file is invalid or is not UTF-8 encoded.") from error
    if payload in (None, "", [], {}):
        raise ValueError("The JSON file does not contain analyzable content.")
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str)


def _analyze_table(headers: list[str], rows: list[list[Any]]) -> dict[str, Any]:
    classified = [_classify_column(_normalize_header(header)) for header in headers]
    revenue_total = 0.0
    expense_total = 0.0
    profit_total = 0.0
    has_profit_column = "profit" in classified
    series_map: dict[str, dict[str, float]] = {}
    category_map: dict[str, float] = {}
    location_map: dict[str, dict[str, Any]] = {}
    numeric_columns: dict[int, float] = {}
    normalized_headers = [_normalize_header(header) for header in headers]

    for row in rows:
        period = None
        category = None
        location = None
        latitude = None
        longitude = None
        revenue = 0.0
        expense = 0.0
        profit = 0.0
        for index, raw in enumerate(row):
            if index >= len(classified):
                break
            kind = classified[index]
            normalized_header = normalized_headers[index]
            if kind == "period":
                period = str(raw).strip() if raw is not None else None
            elif kind == "category":
                category = str(raw).strip() if raw is not None else None
            elif kind == "location":
                location = str(raw).strip() if raw is not None else None
            elif kind == "revenue":
                parsed = _parse_number(raw) or 0.0
                revenue += parsed
            elif kind == "expense":
                parsed = _parse_number(raw) or 0.0
                expense += parsed
            elif kind == "profit":
                parsed = _parse_number(raw) or 0.0
                profit += parsed
            elif normalized_header == "latitude":
                latitude = _parse_number(raw)
            elif normalized_header == "longitude":
                longitude = _parse_number(raw)
            else:
                parsed = _parse_number(raw)
                if parsed is not None:
                    numeric_columns[index] = numeric_columns.get(index, 0.0) + parsed

        if not has_profit_column:
            profit = revenue - expense
        revenue_total += revenue
        expense_total += expense
        profit_total += profit

        label = period or "Total"
        bucket = series_map.setdefault(label, {"revenue": 0.0, "expense": 0.0, "profit": 0.0})
        bucket["revenue"] += revenue
        bucket["expense"] += expense
        bucket["profit"] += profit

        if category:
            category_map[category] = category_map.get(category, 0.0) + (revenue or profit)
        if location:
            location_bucket = location_map.setdefault(location, {"revenue": 0.0, "expense": 0.0, "profit": 0.0})
            location_bucket["revenue"] += revenue
            location_bucket["expense"] += expense
            location_bucket["profit"] += profit
            if latitude is not None:
                location_bucket["latitude"] = latitude
            if longitude is not None:
                location_bucket["longitude"] = longitude

    if revenue_total == 0 and expense_total == 0 and numeric_columns:
        ranked = sorted(numeric_columns.items(), key=lambda item: item[1], reverse=True)
        revenue_total = ranked[0][1]
        if len(ranked) > 1:
            expense_total = ranked[1][1]
            profit_total = revenue_total - expense_total
        else:
            profit_total = revenue_total
        series_map["Total"] = {
            "revenue": revenue_total,
            "expense": expense_total,
            "profit": profit_total,
        }

    series = [
        {"month": label, "revenue": round(values["revenue"], 2), "expense": round(values["expense"], 2), "profit": round(values["profit"], 2)}
        for label, values in series_map.items()
        if values["revenue"] or values["expense"] or values["profit"]
    ]
    categories = [
        {"category": name, "sales": round(value, 2)}
        for name, value in sorted(category_map.items(), key=lambda item: item[1], reverse=True)
    ]
    locations = []
    for name, values in sorted(
        location_map.items(),
        key=lambda item: abs(item[1]["revenue"] or item[1]["profit"]),
        reverse=True,
    ):
        location = {
            "name": name,
            "revenue": round(values["revenue"], 2),
            "expense": round(values["expense"], 2),
            "profit": round(values["profit"], 2),
        }
        if "latitude" in values:
            location["latitude"] = values["latitude"]
        if "longitude" in values:
            location["longitude"] = values["longitude"]
        locations.append(location)
    alerts = []
    if expense_total > revenue_total > 0:
        alerts.append({"label": "Expense overrun", "status": "Needs review"})
    if categories:
        top = categories[0]
        alerts.append({"label": f"{top['category']} concentration", "status": "Monitoring"})

    return {
        "row_count": len(rows),
        "columns": headers,
        "metrics": {
            "revenue": round(revenue_total, 2),
            "expenses": round(expense_total, 2),
            "profit": round(profit_total, 2),
        },
        "series": series,
        "categories": categories,
        "locations": locations,
        "alerts": alerts,
    }


def _labeled_amount(text: str, keys: set[str]) -> float | None:
    lowered = text.lower()
    for key in sorted(keys, key=len, reverse=True):
        match = re.search(
            rf"\b{re.escape(key)}\b\s*[:=-]?\s*\$?\s*([0-9][0-9,]*(?:\.[0-9]+)?)",
            lowered,
        )
        if match:
            parsed = _parse_number(match.group(1))
            if parsed is not None:
                return parsed
    return None


def _metrics_from_text(text: str) -> dict[str, Any] | None:
    revenue = _labeled_amount(text, REVENUE_KEYS)
    expenses = _labeled_amount(text, EXPENSE_KEYS)
    profit = _labeled_amount(text, PROFIT_KEYS)
    if revenue is None and expenses is None and profit is None:
        return None
    revenue_total = revenue or 0.0
    expense_total = expenses or 0.0
    profit_total = profit if profit is not None else revenue_total - expense_total
    return {
        "metrics": {
            "revenue": round(revenue_total, 2),
            "expenses": round(expense_total, 2),
            "profit": round(profit_total, 2),
        }
    }


def process_document(filename: str, content: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {extension or 'unknown'}")

    excerpt = ""
    analysis: dict[str, Any] = {
        "row_count": 0,
        "columns": [],
        "metrics": {"revenue": 0.0, "expenses": 0.0, "profit": 0.0},
        "series": [],
        "categories": [],
        "locations": [],
        "alerts": [],
    }
    status = "Processed"
    logger.info("[PARSE] Start filename=%s type=%s size_bytes=%d", filename, extension, len(content))

    if extension == ".csv":
        headers, rows = _rows_from_csv(content)
        analysis = _analyze_table(headers, rows)
        excerpt = "\n".join([",".join(headers)] + [",".join(str(cell) for cell in row) for row in rows[:25]])
        status = "Analyzed" if analysis["metrics"]["revenue"] or analysis["row_count"] else "Processed"
    elif extension == ".xlsx":
        headers, rows = _rows_from_xlsx(content)
        analysis = _analyze_table(headers, rows)
        excerpt = "\n".join(
            [",".join(headers)] + [",".join("" if cell is None else str(cell) for cell in row) for row in rows[:25]]
        )
        status = "Analyzed" if analysis["metrics"]["revenue"] or analysis["row_count"] else "Processed"
    elif extension == ".xls":
        raise ValueError("Legacy .xls files are not supported. Save the workbook as .xlsx or CSV.")
    elif extension == ".pdf":
        excerpt = _extract_pdf_text(content)
        status = "Processed" if excerpt else "Uploaded"
    elif extension == ".docx":
        excerpt = _extract_docx_text(content)
        status = "Processed" if excerpt else "Uploaded"
    elif extension == ".txt":
        table = _rows_from_text_table(content)
        if table:
            headers, rows = table
            analysis = _analyze_table(headers, rows)
            excerpt = "\n".join(
                [",".join(headers)] + [",".join(str(cell) for cell in row) for row in rows[:25]]
            )
            status = "Analyzed" if analysis["metrics"]["revenue"] or analysis["row_count"] else "Processed"
        else:
            excerpt = content.decode("utf-8", errors="replace")
            status = "Processed"
    elif extension == ".json":
        excerpt = _extract_json_text(content)
        status = "Processed"
    elif extension in IMAGE_EXTENSIONS:
        status = "Uploaded"
        excerpt = f"Image file stored: {filename}"

    text_metrics = _metrics_from_text(excerpt)
    if text_metrics and not (analysis["metrics"]["revenue"] or analysis["metrics"]["expenses"]):
        analysis.update(text_metrics)
        if analysis["metrics"]["revenue"] or analysis["metrics"]["expenses"] or analysis["metrics"]["profit"]:
            status = "Analyzed"

    if extension not in IMAGE_EXTENSIONS and not excerpt.strip():
        raise ValueError("No readable content could be extracted from this file.")

    logger.info(
        "[PARSE] Complete filename=%s rows=%d columns=%d extracted_chars=%d",
        filename,
        analysis["row_count"],
        len(analysis["columns"]),
        len(excerpt),
    )

    return {
        "status": status,
        "type": extension.lstrip(".").upper(),
        "excerpt": excerpt[:12000],
        "processed_at": datetime.now(timezone.utc).isoformat(),
        **analysis,
    }
